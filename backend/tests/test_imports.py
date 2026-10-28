import csv
import io
import pytest
from sqlalchemy import select, func
from openpyxl import Workbook
from app.models import Lead, Property, ImportRow, LeadProperty
from app.schemas import MappingIn, ImportApproval
from app.services.imports import upload_batch, preview_batch, approve_batch, inspect_file
from app.services.normalization import parse_datetime, csv_bytes, guess_mapping

SIMPLE = b"First Name,Last Name,Phone,Property Address,City,State,ZIP,DNC\nJane,Doe,(202) 555-0199,12 Example St,Austin,Texas,78701,no\n"


def preview(db, settings, content=SIMPLE, filename="leads.csv"):
    batch, repeated = upload_batch(db, content, filename, settings)
    mapping = next(iter(batch.sheets.values()))["suggested_mapping"]
    preview_batch(db, batch, MappingIn(sheets=list(batch.sheets), mapping=mapping), settings)
    rows = db.scalars(select(ImportRow).where(ImportRow.batch_id == batch.id)).all()
    return batch, rows


def approve(db, batch, ids, resolutions=None):
    return approve_batch(
        db, batch, ImportApproval(row_ids=ids, resolutions=resolutions or {}, confirmation="IMPORT LOCAL")
    )


def test_csv_upload_mapping_and_normalization(db, settings):
    batch, rows = preview(db, settings)
    assert rows[0].normalized["primary_phone"] == "+12025550199"
    assert rows[0].normalized["state"] == "TX"
    assert rows[0].normalized["full_name"] == "Jane Doe"
    assert rows[0].original["Phone"] == "(202) 555-0199"
    assert not rows[0].errors and batch.counts["new"] == 1
    approve(db, batch, [rows[0].id])
    lead = db.scalar(select(Lead))
    assert lead.original_created_at is None  # Do not fabricate the import date as the source date.
    assert lead.source_do_not_contact is False


def test_csv_delimiter_bom_and_cp1252(settings):
    sheets = inspect_file(
        "Name;Email\nJos\xe9;test@example.test\n".encode("cp1252"), "old.csv", settings.max_import_rows
    )
    assert sheets["CSV"][1][0][1]["Name"] == "José"


def test_xlsx_multiple_sheets_and_original_dates(db, settings):
    wb = Workbook()
    ws = wb.active
    ws.title = "First"
    ws.append(["Full Name", "Email", "Original Lead Creation Date"])
    ws.append(["One Seller", "ONE@example.test", "2015-05-12"])
    ws = wb.create_sheet("Second")
    ws.append(["Full Name", "Email", "Original Lead Creation Date"])
    ws.append(["Two Seller", "TWO@example.test", "2016-05-12"])
    output = io.BytesIO()
    wb.save(output)
    batch, rows = preview(db, settings, output.getvalue(), "old.xlsx")
    assert len(batch.sheets) == 2 and len(rows) == 2
    approve(db, batch, [r.id for r in rows])
    assert {lead.original_created_at.year for lead in db.scalars(select(Lead))} == {2015, 2016}


def test_ambiguous_dates_rejected_unless_explicit():
    with pytest.raises(ValueError, match="Ambiguous date"):
        parse_datetime("03/04/2020")
    assert parse_datetime("03/04/2020", date_order="MDY").month == 3
    assert parse_datetime("03/04/2020", date_order="DMY").month == 4


@pytest.mark.parametrize("date", ["02/30/2020", "not a date", "March 4"])
def test_invalid_dates(date):
    with pytest.raises(ValueError):
        parse_datetime(date)


def test_missing_required_fields(db, settings):
    batch, rows = preview(db, settings, b"Phone,Notes\n2025550199,No seller name\n")
    assert rows[0].state == "invalid" and any("name" in error for error in rows[0].errors)
    with pytest.raises(ValueError, match="invalid"):
        approve(db, batch, [rows[0].id])


def test_mapping_template_reuse(db, settings):
    batch, _ = upload_batch(db, SIMPLE, "legacy.csv", settings)
    request = MappingIn(sheets=["CSV"], mapping=guess_mapping(batch.sheets["CSV"]["headers"]), template_name="Legacy")
    preview_batch(db, batch, request, settings)
    from app.models import MappingTemplate

    assert db.scalar(select(MappingTemplate)).name == "Legacy"


def test_bad_mapping_not_silently_accepted(db, settings):
    batch, _ = upload_batch(db, SIMPLE, "legacy.csv", settings)
    with pytest.raises(ValueError, match="missing source column"):
        preview_batch(db, batch, MappingIn(sheets=["CSV"], mapping={"does_not_exist": "first_name"}), settings)
    with pytest.raises(ValueError, match="at most once"):
        preview_batch(
            db,
            batch,
            MappingIn(sheets=["CSV"], mapping={"First Name": "full_name", "Last Name": "full_name"}),
            settings,
        )


def test_repeat_file_and_reapprove_are_idempotent(db, settings):
    batch, rows = preview(db, settings)
    approve(db, batch, [rows[0].id])
    existing, repeated = upload_batch(db, SIMPLE, "different-name.csv", settings)
    assert repeated and existing.id == batch.id
    approve(db, batch, [rows[0].id])
    assert db.scalar(select(func.count()).select_from(Lead)) == 1


def test_exact_duplicate_cross_file_requires_review(db, settings):
    first, rows = preview(db, settings)
    approve(db, first, [rows[0].id])
    batch, rows = preview(db, settings, SIMPLE + b"\n", "second.csv")
    assert rows[0].duplicate_kind == "exact_duplicate"
    with pytest.raises(ValueError, match="resolve duplicate"):
        approve(db, batch, [rows[0].id])
    approve(db, batch, [rows[0].id], {str(rows[0].id): "skip"})
    assert rows[0].state == "skipped"
    assert db.scalar(select(func.count()).select_from(Lead)) == 1


def test_ambiguous_same_seller_different_property_link(db, settings):
    first, rows = preview(db, settings)
    approve(db, first, [rows[0].id])
    second = SIMPLE.replace(b"12 Example St", b"14 Other St")
    batch, rows = preview(db, settings, second, "second.csv")
    assert rows[0].duplicate_kind == "same_seller_different_property"
    approve(db, batch, [rows[0].id], {str(rows[0].id): "link"})
    assert rows[0].state == "linked"
    assert db.scalar(select(func.count()).select_from(Lead)) == 1
    assert db.scalar(select(func.count()).select_from(LeadProperty)) == 2


def test_same_property_different_seller_preserves_relationship(db, settings):
    first, rows = preview(db, settings)
    approve(db, first, [rows[0].id])
    batch, rows = preview(db, settings, SIMPLE.replace(b"Jane,Doe,(202) 555-0199", b"Other,Seller,(202) 555-0198"))
    assert rows[0].duplicate_kind == "same_property_different_seller"
    approve(db, batch, [rows[0].id], {str(rows[0].id): "separate"})
    assert db.scalar(select(func.count()).select_from(Lead)) == 2
    assert db.scalar(select(func.count()).select_from(Property)) == 1


def test_duplicates_inside_workbook_flagged(db, settings):
    data = SIMPLE + SIMPLE.splitlines()[1] + b"\n"
    batch, rows = preview(db, settings, data)
    assert rows[1].duplicate_kind == "exact_duplicate"
    assert rows[1].candidates[0]["row_id"] == rows[0].id
    approve(db, batch, [r.id for r in rows], {str(rows[1].id): "skip"})
    assert batch.counts["imported"] == 1 and batch.counts["skipped"] == 1


def test_explicit_confirmation_required(db, settings):
    batch, rows = preview(db, settings)
    with pytest.raises(ValueError, match="IMPORT LOCAL"):
        approve_batch(db, batch, ImportApproval(row_ids=[rows[0].id], confirmation="yes"))


def test_interrupted_local_import_can_retry(db, settings, monkeypatch):
    from app.services import imports

    batch, rows = preview(db, settings)
    original = imports.create_local_lead

    def fail(*a, **kw):
        raise RuntimeError("simulated disk issue")

    monkeypatch.setattr(imports, "create_local_lead", fail)
    approve(db, batch, [rows[0].id])
    assert rows[0].state == "failed" and batch.state == "partial"
    monkeypatch.setattr(imports, "create_local_lead", original)
    approve(db, batch, [rows[0].id])
    assert rows[0].state == "imported"
    assert db.scalar(select(func.count()).select_from(Lead)) == 1


def test_formula_injection_protection():
    for dangerous in ['=HYPERLINK("evil")', "+cmd", "-1+2", "@SUM(1)", "   =1", "\tmalicious"]:
        result = csv_bytes([{"value": dangerous}], ["value"]).decode("utf-8-sig")
        value = list(csv.reader(io.StringIO(result)))[1][0]
        assert value.startswith("'")
    assert "'Alice" not in csv_bytes([{"value": "Alice"}], ["value"]).decode()


def test_rejects_duplicate_headers_and_oversized_rows(settings):
    with pytest.raises(ValueError, match="duplicate headers"):
        inspect_file(b"Name,Name\nOne,Two\n", "a.csv", 10)
    with pytest.raises(ValueError, match="exceeds"):
        inspect_file(b"Name,Email\nOne,one@example.test\nTwo,two@example.test\n", "a.csv", 1)


def test_thousands_of_records_preview_and_import(db, settings):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Full Name", "Email", "Property Address", "City", "State"])
    for i in range(2000):
        writer.writerow([f"Fictional Seller {i}", f"test{i}@example.test", f"{i} Sample Street", "Austin", "TX"])
    batch, rows = preview(db, settings, output.getvalue().encode())
    assert batch.counts["new"] == 2000
    approve(db, batch, [r.id for r in rows])
    assert batch.counts["imported"] == 2000
    assert db.scalar(select(func.count()).select_from(Lead)) == 2000


def test_subset_import_stays_partial_until_unresolved_rows_are_reviewed(db, settings):
    data = SIMPLE + b"Other,Seller,2025550198,14 Example St,Austin,TX,78701,no\n"
    batch, rows = preview(db, settings, data)
    approve(db, batch, [rows[0].id])
    assert batch.state == "partial"
    approve(db, batch, [rows[1].id])
    assert batch.state == "completed"


def test_stale_preview_duplicate_is_reported_with_new_candidates(db, settings, make_lead):
    batch, rows = preview(db, settings)
    make_lead(name="Existing Seller", phone="+12025550199")
    approve(db, batch, [rows[0].id])
    assert rows[0].state == "failed" and rows[0].duplicate_kind != "new_record"
    assert rows[0].candidates


def test_xlsx_never_silently_discards_values_without_headers():
    from app.services.imports import inspect_file

    wb = Workbook()
    ws = wb.active
    ws.append(["Full Name", "Phone", None])
    ws.append(["Example Seller", "2025550199", "Unlabelled factual data"])
    content = io.BytesIO()
    wb.save(content)
    with pytest.raises(ValueError, match="without a column header"):
        inspect_file(content.getvalue(), "missing-header.xlsx", 100)


def test_reimpli_epoch_timestamps_become_real_dates():
    """REsimpli sends instants as numbers; every stored date used to be dropped."""
    from datetime import timezone as tz

    from app.services.normalization import parse_epoch

    assert parse_datetime(1761600000000) == parse_datetime("1761600000000")
    assert parse_datetime(1761600000) == parse_datetime(1761600000000)
    assert parse_datetime(1761600000000).tzinfo == tz.utc
    assert parse_datetime(1761600000000).year == 2025
    # Milliseconds and seconds must not be confused for one another.
    assert parse_datetime(1761600000123).microsecond == 123000
    # Non-timestamps are still refused loudly rather than guessed.
    for value in (0, 12345, "not-a-date", 1.5):
        assert parse_epoch(value) is None
